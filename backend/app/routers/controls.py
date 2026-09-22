import csv
import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.database import get_db
from app.models.control import Control
from app.dependencies import get_current_user, require_auditor

router = APIRouter(
    prefix="/api/controls",
    tags=["Controls"],
)


# =========================================================
# GET ALL CONTROLS
# =========================================================

@router.get("")
def get_controls(
    current_user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    controls = db.query(Control).all()

    return controls


# =========================================================
# GET SINGLE CONTROL
# =========================================================

@router.get("/{control_id}")
def get_control(
    control_id: str,
    current_user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    control = (
        db.query(Control)
        .filter(Control.control_id == control_id)
        .first()
    )

    if not control:
        raise HTTPException(
            status_code=404,
            detail="Control not found",
        )

    return control


# =========================================================
# CREATE SINGLE CONTROL
# =========================================================

@router.post("")
def create_control(
    control_id: str,
    title: str,
    framework: str,
    description: str = "",
    requirement: str = "",
    evidence_requested: str = "",
    status: str = "Pending",
    domain: str = "",
    current_user=Depends(require_auditor),
    db: Session = Depends(get_db),
):
    existing_control = (
        db.query(Control)
        .filter(Control.control_id == control_id)
        .first()
    )

    if existing_control:
        raise HTTPException(
            status_code=400,
            detail="Control already exists",
        )

    control = Control(
        control_id=control_id,
        domain=domain,
        title=title,
        framework=framework,
        description=description,
        requirement=requirement,
        evidence_requested=evidence_requested,
        status=status,
        evidence_count=0,
    )

    db.add(control)
    db.commit()
    db.refresh(control)

    return control


# =========================================================
# BULK IMPORT CONTROLS
# CSV / XLSX
# =========================================================

@router.post("/import")
async def import_controls(
    file: UploadFile = File(...),
    current_user=Depends(require_auditor),
    db: Session = Depends(get_db),
):
    # -----------------------------------------------------
    # CHECK FILE
    # -----------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected",
        )

    filename = file.filename.lower()

    if not (
        filename.endswith(".csv")
        or filename.endswith(".xlsx")
    ):
        raise HTTPException(
            status_code=400,
            detail="Only CSV and XLSX files are supported",
        )

    file_content = await file.read()

    try:

        # =================================================
        # READ CSV
        # =================================================

        if filename.endswith(".csv"):

            text = file_content.decode("utf-8-sig")

            reader = csv.DictReader(
                io.StringIO(text)
            )

            rows = list(reader)

        # =================================================
        # READ EXCEL
        # =================================================

        else:

            workbook = load_workbook(
                io.BytesIO(file_content),
                read_only=True,
                data_only=True,
            )

            worksheet = workbook.active

            values = list(
                worksheet.iter_rows(
                    values_only=True
                )
            )

            if not values:
                raise HTTPException(
                    status_code=400,
                    detail="Excel file is empty",
                )

            headers = [
                str(header).strip()
                if header is not None
                else ""
                for header in values[0]
            ]

            rows = [
                dict(zip(headers, row))
                for row in values[1:]
            ]

        # =================================================
        # EMPTY FILE
        # =================================================

        if not rows:
            raise HTTPException(
                status_code=400,
                detail="File contains no data rows",
            )

        # =================================================
        # REQUIRED COLUMNS
        # =================================================

        required_columns = {
    "domain",
    "control_id",
    "control_title",
    "control_description",
    "requirement",
    "status",
    "evidence_requested",
    "framework",
}

        actual_columns = {
            str(column).strip()
            for column in rows[0].keys()
        }

        missing_columns = (
            required_columns - actual_columns
        )

        if missing_columns:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "Missing required columns",
                    "missing_columns": sorted(
                        missing_columns
                    ),
                },
            )

        # =================================================
        # IMPORT VARIABLES
        # =================================================

        inserted = 0
        skipped = 0
        errors = []

        valid_statuses = {
            "Pending",
            "Compliant",
            "Partial",
            "Non-Compliant",
        }

        # =================================================
        # PROCESS EACH ROW
        # =================================================

        for row_number, row in enumerate(
            rows,
            start=2,
        ):

            # -------------------------------------------------
            # READ VALUES
            # -------------------------------------------------

            domain = str(
                row.get("domain") or ""
            ).strip()

            control_id = str(
                row.get("control_id") or ""
            ).strip()

            title = str(
                row.get("control_title") or ""
            ).strip()

            description = str(
                row.get("control_description") or ""
            ).strip()

            requirement = str(
    row.get("requirement") or ""
).strip()
            

            status = str(
                row.get("status") or "Pending"
            ).strip()

            evidence_requested = str(
                row.get("evidence_requested") or ""
            ).strip()

            framework = str(
                row.get("framework") or ""
            ).strip()

            # -------------------------------------------------
            # VALIDATION
            # -------------------------------------------------

            if not control_id:
                errors.append({
                    "row": row_number,
                    "control_id": "",
                    "error": "Control ID is required",
                })
                continue

            if not title:
                errors.append({
                    "row": row_number,
                    "control_id": control_id,
                    "error": "Control title is required",
                })
                continue

            if not framework:
                errors.append({
                    "row": row_number,
                    "control_id": control_id,
                    "error": "Framework is required",
                })
                continue

            if status not in valid_statuses:
                errors.append({
                    "row": row_number,
                    "control_id": control_id,
                    "error": (
                        f"Invalid status: {status}. "
                        f"Allowed values: "
                        f"{', '.join(sorted(valid_statuses))}"
                    ),
                })
                continue

            # -------------------------------------------------
            # CHECK DUPLICATE IN DATABASE
            # -------------------------------------------------

            existing_control = (
                db.query(Control)
                .filter(
                    Control.control_id == control_id
                )
                .first()
            )

            if existing_control:
                skipped += 1
                continue

            # -------------------------------------------------
            # CREATE CONTROL
            # -------------------------------------------------

            control = Control(
    domain=domain,
    control_id=control_id,
    title=title,
    framework=framework,
    description=description,
    requirement=requirement,
    evidence_requested=evidence_requested,
    status=status,
    evidence_count=0,
)

            db.add(control)

            inserted += 1

        # =================================================
        # SAVE TO DATABASE
        # =================================================

        db.commit()

        # =================================================
        # RESPONSE
        # =================================================

        return {
            "message": "Controls imported successfully",
            "inserted": inserted,
            "skipped": skipped,
            "errors": errors,
            "total_rows": len(rows),
        }

    except HTTPException:
        raise

    except Exception as error:

        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=f"Import failed: {str(error)}",
        )



# =========================================================
# UPDATE / EDIT CONTROL
# =========================================================

@router.put("/{control_id}")
def update_control(
    control_id: str,
    title: str = None,
    framework: str = None,
    description: str = None,
    requirement: str = None,
    evidence_requested: str = None,
    status: str = None,
    domain: str = None,
    current_user=Depends(require_auditor),
    db: Session = Depends(get_db),
):
    control = (
        db.query(Control)
        .filter(Control.control_id == control_id)
        .first()
    )

    if not control:
        raise HTTPException(
            status_code=404,
            detail="Control not found",
        )

    if title is not None:
        control.title = title

    if framework is not None:
        control.framework = framework

    if description is not None:
        control.description = description

    if requirement is not None:
        control.requirement = requirement

    if evidence_requested is not None:
        control.evidence_requested = evidence_requested

    if status is not None:
        valid_statuses = {
            "Pending",
            "Compliant",
            "Partial",
            "Non-Compliant",
        }

        if status not in valid_statuses:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Invalid status: {status}. "
                    f"Allowed values: {', '.join(sorted(valid_statuses))}"
                ),
            )

        control.status = status

    if domain is not None:
        control.domain = domain

    db.commit()
    db.refresh(control)

    return control